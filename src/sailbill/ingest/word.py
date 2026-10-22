"""Word (.docx) ingestion: paragraphs + tables, merged cells de-duplicated."""

from __future__ import annotations

from pathlib import Path

from ..models import NormalizedDocument, SourceKind, TableData


def parse_word(doc: NormalizedDocument) -> None:
    path = Path(doc.path) if doc.path else None
    if path is None:
        return
    if path.suffix.lower() == ".doc":
        doc.warnings.append("旧版 .doc 需要先转换为 .docx（soffice --convert-to docx）")
        doc.meta["needs_conversion"] = True
        return

    try:
        import docx  # python-docx
    except ImportError:  # pragma: no cover
        doc.warnings.append("未安装 python-docx，无法解析 Word")
        return

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        doc.warnings.append(f"Word 打开失败: {exc}")
        return

    for i, para in enumerate(document.paragraphs):
        text = (para.text or "").strip()
        if text:
            style = (getattr(para.style, "name", "") or "").lower()
            kind = "heading" if "heading" in style or "标题" in style else "paragraph"
            doc.add_block(text, kind=kind, line=i, source=doc.filename)

    for t, table in enumerate(document.tables):
        rows: list[list[str]] = []
        for row in table.rows:
            cells: list[str] = []
            previous_tc = None
            for cell in row.cells:
                # a merged cell is exposed once per grid column; detect it by
                # identity, never by equal text (two real cells may both be 40.00)
                if cell._tc is previous_tc:
                    continue
                previous_tc = cell._tc
                cells.append((cell.text or "").strip())
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        header_row = _guess_header_row(rows)
        table_data = TableData(
            table_id=f"{doc.doc_id}-t{t + 1}",
            rows=rows,
            header_row=header_row,
            headers=rows[header_row] if header_row is not None else [],
            source=doc.filename,
        )
        doc.tables.append(table_data)
        for i, row in enumerate(rows):
            doc.add_block(
                " | ".join(v for v in row if v),
                kind="table_cell",
                table_id=table_data.table_id,
                row=i,
                source=doc.filename,
            )

    doc.text = "\n".join(b.text for b in doc.blocks)
    if not doc.blocks:
        doc.warnings.append("Word 中没有可提取的文本")


def _guess_header_row(rows: list[list[str]], window: int = 10) -> int | None:
    best, best_score = None, 0.0
    for i, row in enumerate(rows[:window]):
        cells = [c for c in row if c]
        if len(cells) < 2:
            continue
        numeric = sum(1 for c in cells if _numeric(c))
        score = (len(cells) - numeric) - numeric * 1.5
        if score > best_score:
            best, best_score = i, score
    return best


def _numeric(text: str) -> bool:
    t = text.replace(",", "").replace(".", "").replace("-", "").strip()
    return bool(t) and t.isdigit()


def kind_of(path: str) -> SourceKind:
    return SourceKind.WORD
